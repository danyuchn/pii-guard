"""Tests for reference lists: inspection, shape rules and stored sources.

Every value in this file is invented.  The tests that matter most are the ones
asserting that a value never leaves the module: the whole point of a reference
list is that the table stays where it is.
"""

from __future__ import annotations

import json
import stat
from pathlib import Path

import pytest

from pii_guard import reference
from pii_guard.local_workflow import WorkflowError

pytest.importorskip("openpyxl")

HEADER = ["姓名", "公司", "手機", "電子郵件", "訂單編號", "金額"]
ROWS = [
    ["陳建華", "寶島顧問有限公司", "0912345678", "chen@example.invalid", "ORD-000101", "12000"],
    ["林雅婷", "海風科技股份有限公司", "0933112244", "lin@example.invalid", "ORD-000102", "8500"],
    ["黃國強", "松風設計工作室", "0955667788", "huang@example.invalid", "ORD-000103", "23000"],
    ["李", "阿發企業社", "0922334455", "lee@example.invalid", "ORD-000104", "500"],
]


def write_xlsx(path: Path, *, sheet: str = "客戶") -> Path:
    import openpyxl

    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = sheet
    worksheet.append(HEADER)
    for row in ROWS:
        worksheet.append(row)
    workbook.save(path)
    return path


def write_csv(path: Path, encoding: str = "utf-8") -> Path:
    lines = [",".join(HEADER)] + [",".join(row) for row in ROWS]
    path.write_bytes(("\n".join(lines) + "\n").encode(encoding))
    return path


def column(report: reference.TableReport, name: str) -> reference.ColumnReport:
    for entry in report.columns:
        if entry.name == name:
            return entry
    raise AssertionError(f"no column named {name}")


# ---------------------------------------------------------------------------
# Header and value guessing
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("header", "expected"),
    [
        ("客戶姓名", "PERSON"),
        ("Contact Name", "PERSON"),
        ("公司抬頭", "ORG"),
        ("手機號碼", "TW_MOBILE"),
        ("市話", "TW_LANDLINE"),
        ("Email", "EMAIL_ADDRESS"),
        ("統一編號", "TW_BUSINESS_ID"),
        ("身分證字號", "TW_NATIONAL_ID"),
        ("通訊地址", "TW_ADDRESS"),
        ("車牌號碼", "TW_LICENSE_PLATE"),
        ("訂單編號", "ORDER_ID"),
        ("成交金額", "SKIP"),
        ("建立日期", "SKIP"),
    ],
)
def test_header_keywords_name_the_column(header: str, expected: str) -> None:
    assert reference.guess_from_header(header) == expected


def test_longer_keywords_win_over_shorter_ones() -> None:
    assert reference.guess_from_header("統一編號") == "TW_BUSINESS_ID"
    assert reference.guess_from_header("案件編號") == "ORDER_ID"


def test_unknown_type_names_fall_back_without_crashing() -> None:
    assert reference.normalize_type("tw-id") == "TW_NATIONAL_ID"
    assert reference.normalize_type("PERSON") == "PERSON"
    assert reference.normalize_type("MY_CODE") == "MY_CODE"
    assert reference.normalize_type("不是類型") == "CUSTOM"


# ---------------------------------------------------------------------------
# Risk rules
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("value", ["李", "abc", "123456", "", "   "])
def test_short_values_are_risky(value: str) -> None:
    assert reference.is_risky(value) is True


@pytest.mark.parametrize("value", ["陳建華", "Chen", "0912345678", "1234567"])
def test_ordinary_values_are_not_risky(value: str) -> None:
    assert reference.is_risky(value) is False


# ---------------------------------------------------------------------------
# Shape inference
# ---------------------------------------------------------------------------


def test_a_consistent_column_yields_one_shape() -> None:
    regex = reference.infer_shape(["ORD-000101", "ORD-000102", "ORD-000987"])

    assert regex is not None
    import re

    assert re.search(regex, "訂單 ORD-000555 已出貨")
    assert not re.search(regex, "ORD-00055")


def test_a_mixed_column_has_no_shape() -> None:
    assert reference.infer_shape(["ORD-000101", "A1", "客戶自取", "2026/01/01"]) is None


def test_a_short_shape_is_refused() -> None:
    # Four characters would match far too much ordinary text.
    assert reference.infer_shape(["A123", "B456", "C789"]) is None


def test_a_column_of_one_repeated_literal_has_no_shape() -> None:
    assert reference.infer_shape(["未結案", "未結案", "未結案"]) is None


def test_a_shape_needs_nine_in_ten_to_agree() -> None:
    clean = [f"ORD-00{index:04d}" for index in range(10)]

    # One stray value in ten still clears the 90% bar.
    assert reference.infer_shape([*clean, "待補"]) is not None
    # Two do not.
    assert reference.infer_shape([*clean, "待補", "客戶自取"]) is None


# ---------------------------------------------------------------------------
# Inspection
# ---------------------------------------------------------------------------


def test_inspect_reads_an_xlsx_and_names_its_columns(tmp_path: Path) -> None:
    report = reference.inspect_table(write_xlsx(tmp_path / "list.xlsx"))

    assert report.rows == len(ROWS)
    assert report.sheet == "客戶"
    assert report.sheets == ("客戶",)
    assert column(report, "姓名").guessed_type == "PERSON"
    assert column(report, "公司").guessed_type == "ORG"
    assert column(report, "手機").guessed_type == "TW_MOBILE"
    assert column(report, "金額").guessed_type == "SKIP"


def test_inspect_reads_a_csv(tmp_path: Path) -> None:
    report = reference.inspect_table(write_csv(tmp_path / "list.csv"))

    assert report.rows == len(ROWS)
    assert column(report, "訂單編號").guessed_type == "ORDER_ID"


def test_inspect_reads_a_cp950_csv(tmp_path: Path) -> None:
    report = reference.inspect_table(write_csv(tmp_path / "big5.csv", encoding="cp950"))

    assert column(report, "姓名").guessed_type == "PERSON"
    assert column(report, "姓名").non_empty == len(ROWS)


def test_inspect_counts_the_risky_values(tmp_path: Path) -> None:
    report = reference.inspect_table(write_xlsx(tmp_path / "list.xlsx"))

    # Only the one-character surname is too short to mask safely.
    assert column(report, "姓名").risky == 1
    assert column(report, "手機").risky == 0


def test_inspect_keeps_at_most_three_samples(tmp_path: Path) -> None:
    report = reference.inspect_table(write_xlsx(tmp_path / "list.xlsx"))

    assert len(column(report, "姓名").samples) == reference.SAMPLE_LIMIT


def test_the_describe_payload_withholds_samples_by_default(tmp_path: Path) -> None:
    report = reference.inspect_table(write_xlsx(tmp_path / "list.xlsx"))

    payload = json.dumps(report.describe(), ensure_ascii=False)

    assert "陳建華" not in payload
    assert "0912345678" not in payload
    assert "samples" not in payload


def test_describe_includes_samples_only_when_asked(tmp_path: Path) -> None:
    report = reference.inspect_table(write_xlsx(tmp_path / "list.xlsx"))

    payload = json.dumps(report.describe(include_samples=True), ensure_ascii=False)

    assert "陳建華" in payload


def test_a_missing_table_is_a_clean_failure(tmp_path: Path) -> None:
    with pytest.raises(WorkflowError) as error:
        reference.inspect_table(tmp_path / "nope.xlsx")

    assert error.value.code == "TABLE_NOT_FOUND"


def test_an_unsupported_format_is_refused(tmp_path: Path) -> None:
    target = tmp_path / "list.docx"
    target.write_bytes(b"not a table")

    with pytest.raises(WorkflowError) as error:
        reference.inspect_table(target)

    assert error.value.code == "TABLE_UNSUPPORTED"


def test_a_named_sheet_can_be_chosen(tmp_path: Path) -> None:
    import openpyxl

    path = write_xlsx(tmp_path / "list.xlsx")
    workbook = openpyxl.load_workbook(path)
    second = workbook.create_sheet("備用")
    second.append(["車牌"])
    second.append(["ABC-1234"])
    workbook.save(path)

    report = reference.inspect_table(path, sheet="備用")

    assert report.sheet == "備用"
    assert column(report, "車牌").guessed_type == "TW_LICENSE_PLATE"

    with pytest.raises(WorkflowError) as error:
        reference.inspect_table(path, sheet="不存在")

    assert error.value.code == "SHEET_NOT_FOUND"


# ---------------------------------------------------------------------------
# Stored sources
# ---------------------------------------------------------------------------


def source_for(path: Path) -> reference.ReferenceSource:
    return reference.ReferenceSource(
        path=str(path),
        sheet="客戶",
        columns={"姓名": "PERSON", "手機": "TW_MOBILE", "訂單編號": "ORDER_ID"},
        patterns=(("ORDER_ID", r"(?<![A-Za-z0-9])[A-Z]{3}\-\d{6}(?![A-Za-z0-9])"),),
    )


def test_sources_round_trip(tmp_path: Path) -> None:
    table = write_xlsx(tmp_path / "list.xlsx")
    project = tmp_path / "project"
    project.mkdir()

    reference.write_sources(project, [source_for(table)])
    loaded = reference.load_sources(project)

    assert len(loaded) == 1
    assert loaded[0].columns["姓名"] == "PERSON"
    assert loaded[0].sheet == "客戶"
    assert loaded[0].patterns[0][0] == "ORDER_ID"


def test_the_stored_description_holds_no_values(tmp_path: Path) -> None:
    table = write_xlsx(tmp_path / "list.xlsx")
    project = tmp_path / "project"
    project.mkdir()

    saved = reference.write_sources(project, [source_for(table)])
    text = saved.read_text(encoding="utf-8")

    assert "陳建華" not in text
    assert "0912345678" not in text
    assert "ORD-000101" not in text


def test_the_stored_description_is_owner_only(tmp_path: Path) -> None:
    project = tmp_path / "project"
    project.mkdir()

    saved = reference.write_sources(project, [source_for(tmp_path / "list.xlsx")])

    assert stat.S_IMODE(saved.stat().st_mode) == 0o600
    assert stat.S_IMODE(saved.parent.stat().st_mode) == 0o700


def test_writing_sources_adds_the_directory_to_gitignore(tmp_path: Path) -> None:
    project = tmp_path / "project"
    project.mkdir()

    reference.write_sources(project, [])

    assert ".pii-guard/" in (project / ".gitignore").read_text(encoding="utf-8")


def test_gitignore_is_not_duplicated(tmp_path: Path) -> None:
    project = tmp_path / "project"
    project.mkdir()
    (project / ".gitignore").write_text("build/\n.pii-guard/\n", encoding="utf-8")

    reference.write_sources(project, [])

    assert (project / ".gitignore").read_text(encoding="utf-8").count(".pii-guard/") == 1


def test_gitignore_keeps_a_file_without_a_trailing_newline_readable(tmp_path: Path) -> None:
    project = tmp_path / "project"
    project.mkdir()
    (project / ".gitignore").write_text("build/", encoding="utf-8")

    reference.write_sources(project, [])

    lines = (project / ".gitignore").read_text(encoding="utf-8").splitlines()
    assert lines == ["build/", ".pii-guard/"]


def test_a_malformed_sources_file_reads_as_no_sources(tmp_path: Path) -> None:
    target = tmp_path / "sources.json"
    target.write_text("{not json", encoding="utf-8")

    assert reference.load_sources_file(target) == []


def test_an_uncompilable_stored_pattern_is_dropped(tmp_path: Path) -> None:
    target = tmp_path / "sources.json"
    target.write_text(
        json.dumps(
            {
                "version": 1,
                "sources": [
                    {"path": "/tmp/x.csv", "patterns": [{"type": "ORDER_ID", "regex": "[("}]}
                ],
            }
        ),
        encoding="utf-8",
    )

    assert reference.load_sources_file(target)[0].patterns == ()


# ---------------------------------------------------------------------------
# Turning sources into terms
# ---------------------------------------------------------------------------


def test_terms_come_back_typed_and_deduplicated(tmp_path: Path) -> None:
    table = write_xlsx(tmp_path / "list.xlsx")

    loaded = reference.load_reference_terms([source_for(table)])

    assert ("PERSON", "陳建華") in loaded.terms
    assert ("TW_MOBILE", "0912345678") in loaded.terms
    assert ("ORDER_ID", "ORD-000101") in loaded.terms
    assert len(loaded.terms) == len({value for _, value in loaded.terms})


def test_risky_values_are_skipped_and_counted(tmp_path: Path) -> None:
    table = write_xlsx(tmp_path / "list.xlsx")

    loaded = reference.load_reference_terms([source_for(table)])

    assert ("PERSON", "李") not in loaded.terms
    assert loaded.risky_skipped == 1

    with_risky = reference.load_reference_terms([source_for(table)], include_risky=True)
    assert ("PERSON", "李") in with_risky.terms


def test_columns_the_user_did_not_map_are_left_alone(tmp_path: Path) -> None:
    table = write_xlsx(tmp_path / "list.xlsx")

    loaded = reference.load_reference_terms([source_for(table)])

    assert not any(value == "12000" for _, value in loaded.terms)
    assert not any(value.endswith("example.invalid") for _, value in loaded.terms)


def test_a_missing_source_is_reported_not_raised(tmp_path: Path) -> None:
    loaded = reference.load_reference_terms([source_for(tmp_path / "gone.xlsx")])

    assert loaded.terms == ()
    assert loaded.missing == (str(tmp_path / "gone.xlsx"),)


def test_the_term_ceiling_is_honoured(tmp_path: Path) -> None:
    table = write_xlsx(tmp_path / "list.xlsx")

    loaded = reference.load_reference_terms([source_for(table)], max_terms=2)

    assert len(loaded.terms) == 2
    assert loaded.truncated is True


def test_the_summary_reports_counts_and_no_values(tmp_path: Path) -> None:
    table = write_xlsx(tmp_path / "list.xlsx")

    payload = json.dumps(
        reference.load_reference_terms([source_for(table)]).summary(), ensure_ascii=False
    )

    assert "陳建華" not in payload
    assert "0912345678" not in payload
    assert '"terms"' in payload


def test_materialize_writes_an_owner_only_term_file(tmp_path: Path) -> None:
    table = write_xlsx(tmp_path / "list.xlsx")
    target = tmp_path / "project" / ".pii-guard" / "terms.txt"

    written = reference.materialize([source_for(table)], target)

    assert written > 0
    assert stat.S_IMODE(target.stat().st_mode) == 0o600
    body = target.read_text(encoding="utf-8")
    assert "PERSON\t陳建華" in body


def test_materialized_terms_read_back_through_the_policy_loader(tmp_path: Path) -> None:
    from pii_guard.hookd import policy

    table = write_xlsx(tmp_path / "list.xlsx")
    target = tmp_path / "terms.txt"
    reference.materialize([source_for(table)], target)

    assert ("PERSON", "陳建華") in policy.load_seed_terms([str(target)])


# ---------------------------------------------------------------------------
# The cache behind "edit the spreadsheet and it just works"
# ---------------------------------------------------------------------------


def test_the_cache_reloads_when_the_table_changes(tmp_path: Path) -> None:
    import os
    import openpyxl

    table = write_xlsx(tmp_path / "list.xlsx")
    project = tmp_path / "project"
    project.mkdir()
    reference.write_sources(project, [source_for(table)])
    cache = reference.ReferenceCache([str(reference.sources_path(project))])

    before = cache.load()
    assert not any(value == "吳孟儒" for _, value in before.terms)

    workbook = openpyxl.load_workbook(table)
    workbook["客戶"].append(
        ["吳孟儒", "新芽行銷", "0911222333", "wu@example.invalid", "ORD-000105", "999"]
    )
    workbook.save(table)
    # Two writes inside one filesystem timestamp tick would look unchanged.
    status = table.stat()
    os.utime(table, ns=(status.st_atime_ns, status.st_mtime_ns + 1_000_000_000))

    after = cache.load()

    assert any(value == "吳孟儒" for _, value in after.terms)


def test_the_cache_does_not_reread_an_unchanged_table(tmp_path: Path) -> None:
    table = write_xlsx(tmp_path / "list.xlsx")
    project = tmp_path / "project"
    project.mkdir()
    reference.write_sources(project, [source_for(table)])
    cache = reference.ReferenceCache([str(reference.sources_path(project))])

    assert cache.load() is cache.load()


def test_invalidate_forces_a_reread(tmp_path: Path) -> None:
    table = write_xlsx(tmp_path / "list.xlsx")
    project = tmp_path / "project"
    project.mkdir()
    reference.write_sources(project, [source_for(table)])
    cache = reference.ReferenceCache([str(reference.sources_path(project))])

    first = cache.load()
    cache.invalidate()

    assert cache.load() is not first


# ---------------------------------------------------------------------------
# Leading zeros a spreadsheet dropped
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("entity_type", "stored", "expected"),
    [
        ("TW_MOBILE", "912345678", "0912345678"),
        ("TW_LANDLINE", "223456789", "0223456789"),
        ("TW_LANDLINE", "47654321", "047654321"),
        ("TW_BUSINESS_ID", "4567890", "04567890"),
    ],
)
def test_a_dropped_leading_zero_is_put_back(
    entity_type: str, stored: str, expected: str
) -> None:
    assert reference.normalize_value(entity_type, stored) == expected


@pytest.mark.parametrize(
    ("entity_type", "value"),
    [
        # Already complete: nothing was lost.
        ("TW_MOBILE", "0912345678"),
        # Punctuation means the cell was text, so the zero was never dropped.
        ("TW_LANDLINE", "02-23456789"),
        ("TW_LANDLINE", "0223456789"),
        # A full business id is exactly eight digits.
        ("TW_BUSINESS_ID", "12345678"),
        # A name is never rewritten, whatever it looks like.
        ("PERSON", "912345678"),
        ("ORDER_ID", "912345678"),
        ("CUSTOM", "9123"),
    ],
)
def test_a_complete_value_is_left_alone(entity_type: str, value: str) -> None:
    assert reference.normalize_value(entity_type, value) == value


def test_a_numeric_mobile_column_yields_the_dialled_number(tmp_path: Path) -> None:
    import openpyxl

    path = tmp_path / "numeric.xlsx"
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "客戶"
    worksheet.append(["姓名", "手機", "統一編號"])
    # Written as numbers, exactly as Excel stores them without a text format.
    worksheet.append(["陳建華", 912345678, 4567890])
    workbook.save(path)

    report = reference.inspect_table(path)
    assert column(report, "手機").samples == ("0912345678",)

    source = reference.ReferenceSource(
        path=str(path),
        sheet="客戶",
        columns={"手機": "TW_MOBILE", "統一編號": "TW_BUSINESS_ID"},
    )
    loaded = reference.load_reference_terms([source])

    assert ("TW_MOBILE", "0912345678") in loaded.terms
    assert ("TW_BUSINESS_ID", "04567890") in loaded.terms


def test_materialized_terms_carry_the_restored_zero(tmp_path: Path) -> None:
    import openpyxl

    path = tmp_path / "numeric.xlsx"
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.append(["手機"])
    worksheet.append([912345678])
    workbook.save(path)
    target = tmp_path / "terms.txt"

    reference.materialize(
        [reference.ReferenceSource(path=str(path), columns={"手機": "TW_MOBILE"})], target
    )

    assert "TW_MOBILE\t0912345678" in target.read_text(encoding="utf-8")


# ---------------------------------------------------------------------------
# Registering a project with the service
# ---------------------------------------------------------------------------


@pytest.fixture()
def installer_config(tmp_path: Path, monkeypatch) -> Path:
    target = tmp_path / "config" / "hookd.json"
    monkeypatch.setenv("PII_GUARD_HOOKD_CONFIG", str(target))
    return target


def test_registering_a_project_records_its_description_file(
    tmp_path: Path, installer_config: Path
) -> None:
    project = tmp_path / "project"
    project.mkdir()

    assert reference.register_project(project) is True

    payload = json.loads(installer_config.read_text(encoding="utf-8"))
    assert payload["policy"]["reference_sources"] == [str(reference.sources_path(project))]


def test_registering_twice_changes_nothing(tmp_path: Path, installer_config: Path) -> None:
    project = tmp_path / "project"
    project.mkdir()

    assert reference.register_project(project) is True
    assert reference.register_project(project) is False

    payload = json.loads(installer_config.read_text(encoding="utf-8"))
    assert len(payload["policy"]["reference_sources"]) == 1


def test_registering_keeps_the_rest_of_the_config(
    tmp_path: Path, installer_config: Path
) -> None:
    installer_config.parent.mkdir(parents=True)
    installer_config.write_text(
        json.dumps(
            {
                "repo": "/somewhere/pii-guard",
                "engine": "full",
                "serve_command": ["uv", "run", "pii-guard-hookd", "serve"],
                "policy": {"allowed_tools": ["mcp__local-db__query"], "output_gate": False},
            }
        ),
        encoding="utf-8",
    )
    project = tmp_path / "project"
    project.mkdir()

    reference.register_project(project)

    payload = json.loads(installer_config.read_text(encoding="utf-8"))
    assert payload["serve_command"] == ["uv", "run", "pii-guard-hookd", "serve"]
    assert payload["policy"]["allowed_tools"] == ["mcp__local-db__query"]
    assert payload["policy"]["output_gate"] is False
    assert payload["policy"]["reference_sources"] == [str(reference.sources_path(project))]


def test_the_registered_config_is_owner_only(tmp_path: Path, installer_config: Path) -> None:
    project = tmp_path / "project"
    project.mkdir()

    reference.register_project(project)

    assert stat.S_IMODE(installer_config.stat().st_mode) == 0o600


def test_unregistering_removes_only_that_project(
    tmp_path: Path, installer_config: Path
) -> None:
    first, second = tmp_path / "one", tmp_path / "two"
    first.mkdir()
    second.mkdir()
    reference.register_project(first)
    reference.register_project(second)

    assert reference.unregister_project(first) is True
    assert reference.unregister_project(first) is False

    assert reference.registered_source_files() == (str(reference.sources_path(second)),)


def test_registered_files_are_empty_without_a_config(installer_config: Path) -> None:
    assert reference.registered_source_files() == ()


def test_a_cache_over_the_registry_sees_a_project_registered_later(
    tmp_path: Path, installer_config: Path
) -> None:
    table = write_xlsx(tmp_path / "list.xlsx")
    project = tmp_path / "project"
    project.mkdir()
    cache = reference.ReferenceCache(reference.registered_source_files)

    assert cache.load().terms == ()

    reference.write_sources(project, [source_for(table)])
    reference.register_project(project)

    assert any(value == "陳建華" for _, value in cache.load().terms)


def test_a_cache_over_the_registry_forgets_an_unregistered_project(
    tmp_path: Path, installer_config: Path
) -> None:
    table = write_xlsx(tmp_path / "list.xlsx")
    project = tmp_path / "project"
    project.mkdir()
    reference.write_sources(project, [source_for(table)])
    reference.register_project(project)
    cache = reference.ReferenceCache(reference.registered_source_files)
    assert cache.load().terms != ()

    reference.unregister_project(project)

    assert cache.load().terms == ()
